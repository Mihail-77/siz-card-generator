import base64
import os
import subprocess
import sys
import threading
from datetime import date, datetime, timedelta
from pathlib import Path
import tkinter as tk
from tkinter import messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from siz_card_generator.generator import (
    create_single_card,
    get_departments,
    get_positions_by_department,
)



def get_application_dir():
    """Return the folder containing the executable or the source app."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def get_resource_path(filename):
    """Return a bundled resource path or a path next to the source file."""
    resource_dir = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return resource_dir / filename


BASE_DIR = get_application_dir()
APP_ICON_PATH = get_resource_path("app_icon.ico")
DATA_DIR = BASE_DIR / "data"
NORMS_PATH = DATA_DIR / "norms.xlsx"
TEMPLATE_PATH = BASE_DIR / "templates" / "card_template.xlsx"
OUTPUT_DIR = BASE_DIR / "output"
BACKUPS_DIR = OUTPUT_DIR / "backups"
REQUESTS_DIR = BASE_DIR / "requests"

REQUEST_TITLE = "Заявка на добавление должности в нормы выдачи СИЗ"
REQUEST_FIELDS = [
    "Дата заявки",
    "Подразделение",
    "Должность",
    "Описание выполняемых работ",
    "Основание для добавления в нормы СИЗ",
    "Предполагаемые СИЗ, если известно",
    "ФИО инициатора заявки",
    "Контакт инициатора",
    "Статус рассмотрения",
    "Решение специалиста по ОТ",
    "Комментарий специалиста по ОТ",
]
REQUEST_REQUIRED_FIELDS = [
    "Подразделение",
    "Должность",
    "Описание выполняемых работ",
    "Основание для добавления в нормы СИЗ",
    "ФИО инициатора заявки",
    "Контакт инициатора",
]
REQUEST_MULTILINE_FIELDS = {
    "Описание выполняемых работ",
    "Основание для добавления в нормы СИЗ",
    "Предполагаемые СИЗ, если известно",
}
REQUEST_REVIEW_FIELDS = {
    "Статус рассмотрения",
    "Решение специалиста по ОТ",
    "Комментарий специалиста по ОТ",
}

PRINT_PREVIEW_SCRIPT = r"""
$ErrorActionPreference = "Stop"
[Console]::OutputEncoding = [System.Text.Encoding]::UTF8

$excel = $null
$workbooks = $null
$workbook = $null
$worksheetCollection = $null
$worksheets = @()
$errorMessage = $null

try {
    $cardPath = $env:SIZ_CARD_PATH
    if ([string]::IsNullOrWhiteSpace($cardPath) -or -not (Test-Path -LiteralPath $cardPath -PathType Leaf)) {
        throw "Файл карточки не найден."
    }

    try {
        $excel = New-Object -ComObject Excel.Application
    }
    catch {
        throw "Для предварительного просмотра и печати требуется установленный Microsoft Excel. Создание карточек доступно без Excel."
    }
    $excel.DisplayAlerts = $false
    $workbooks = $excel.Workbooks
    $workbook = $workbooks.Open($cardPath)
    $worksheetCollection = $workbook.Worksheets

    $sheetSettings = @{
        "Лицевая сторона" = @{
            Left = 0.7874015748031497
            Right = 0.3149606299212598
        }
        "Оборотная сторона" = @{
            Left = 0.3937007874015748
            Right = 0.7086614173228347
        }
    }

    $totalPages = 0
    foreach ($sheetName in @($sheetSettings.Keys)) {
        $worksheet = $null
        $pageSetup = $null
        $pages = $null
        try {
            $worksheet = $worksheetCollection.Item($sheetName)
        }
        catch {
            throw "В карточке нет обязательного листа: $sheetName."
        }
        $worksheets += $worksheet

        $pageSetup = $worksheet.PageSetup
        try {
            if ([string]::IsNullOrWhiteSpace([string]$pageSetup.PrintArea)) {
                $existingArea = $null
                $names = $null
                try {
                    $names = $worksheet.Names
                    foreach ($definedName in @($names)) {
                        try {
                            if ([string]$definedName.Name -like "*!Print_Area") {
                                $reference = [string]$definedName.RefersTo
                                $separator = $reference.IndexOf("!")
                                if ($separator -ge 0) {
                                    $existingArea = $reference.Substring($separator + 1)
                                }
                            }
                        }
                        finally {
                            [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($definedName)
                        }
                    }
                }
                finally {
                    if ($names) {
                        [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($names)
                    }
                }
                if ([string]::IsNullOrWhiteSpace($existingArea)) {
                    throw "Для листа '$sheetName' не определена область печати."
                }
                $pageSetup.PrintArea = $existingArea
            }

            if ($pageSetup.PaperSize -ne 9) { $pageSetup.PaperSize = 9 }
            if ($pageSetup.Orientation -ne 1) { $pageSetup.Orientation = 1 }
            if ([bool]$pageSetup.Zoom) { $pageSetup.Zoom = $false }
            if ($pageSetup.FitToPagesWide -ne 1) { $pageSetup.FitToPagesWide = 1 }
            if ($pageSetup.FitToPagesTall -ne 1) { $pageSetup.FitToPagesTall = 1 }

            $settings = $sheetSettings[$sheetName]
            $leftMargin = $excel.InchesToPoints($settings.Left)
            $rightMargin = $excel.InchesToPoints($settings.Right)
            $topMargin = $excel.InchesToPoints(0.5905511811023623)
            $bottomMargin = $excel.InchesToPoints(0.3937007874015748)
            $headerFooterMargin = $excel.InchesToPoints(0.1968503937007874)
            if ([Math]::Abs($pageSetup.LeftMargin - $leftMargin) -gt 0.001) {
                $pageSetup.LeftMargin = $leftMargin
            }
            if ([Math]::Abs($pageSetup.RightMargin - $rightMargin) -gt 0.001) {
                $pageSetup.RightMargin = $rightMargin
            }
            if ([Math]::Abs($pageSetup.TopMargin - $topMargin) -gt 0.001) {
                $pageSetup.TopMargin = $topMargin
            }
            if ([Math]::Abs($pageSetup.BottomMargin - $bottomMargin) -gt 0.001) {
                $pageSetup.BottomMargin = $bottomMargin
            }
            if ([Math]::Abs($pageSetup.HeaderMargin - $headerFooterMargin) -gt 0.001) {
                $pageSetup.HeaderMargin = $headerFooterMargin
            }
            if ([Math]::Abs($pageSetup.FooterMargin - $headerFooterMargin) -gt 0.001) {
                $pageSetup.FooterMargin = $headerFooterMargin
            }

            $pages = $pageSetup.Pages
            $pageCount = $pages.Count
            if ($pageCount -ne 1) {
                throw "Лист '$sheetName' не помещается на одну печатную страницу."
            }
            $totalPages += $pageCount
        }
        finally {
            if ($pages) {
                [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($pages)
            }
            if ($pageSetup) {
                [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($pageSetup)
            }
        }
    }

    if ($totalPages -ne 2) {
        throw "Предварительный просмотр должен содержать две страницы."
    }

    Write-Output "PREVIEW_PAGES=$totalPages"
    $worksheetCollection.Select()
    $excel.Visible = $true
    $workbook.PrintPreview()
}
catch {
    $errorMessage = $_.Exception.Message
}
finally {
    if ($workbook) {
        try { $workbook.Close($false) } catch {}
    }
    for ($index = $worksheets.Count - 1; $index -ge 0; $index--) {
        if ($worksheets[$index]) {
            try { [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($worksheets[$index]) } catch {}
        }
    }
    if ($worksheetCollection) {
        try { [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($worksheetCollection) } catch {}
    }
    if ($workbook) {
        try { [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbook) } catch {}
    }
    if ($workbooks) {
        try { [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbooks) } catch {}
    }
    if ($excel) {
        try { $excel.Quit() } catch {}
        try { [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($excel) } catch {}
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}

if ($errorMessage) {
    $errorBytes = [System.Text.Encoding]::UTF8.GetBytes($errorMessage)
    Write-Output "ERROR_BASE64=$([Convert]::ToBase64String($errorBytes))"
    exit 1
}
"""


FIELD_LABELS = [
    "Номер карточки",
    "ФИО",
    "Дата приема",
    "Пол",
    "Рост",
    "Размер одежды (верх)",
    "Размер одежды (низ)",
    "Размер обуви",
    "Размер головного убора",
    "Размер перчаток",
]

FIELD_TO_EMPLOYEE_COLUMN = {
    "Номер карточки": "Табельный номер",
}

REQUIRED_FIELDS = [
    "Подразделение",
    "Должность",
    "Номер карточки",
    "ФИО",
    "Дата приема",
    "Пол",
    "Рост",
]

NAVIGATION_KEYS = {
    "Left",
    "Right",
    "Home",
    "End",
    "Tab",
    "Return",
    "Escape",
    "Up",
    "Down",
}


def validate_request_data(request_data):
    return [
        field
        for field in REQUEST_REQUIRED_FIELDS
        if not str(request_data.get(field, "") or "").strip()
    ]


def get_request_path(request_time=None):
    timestamp = (request_time or datetime.now()).replace(microsecond=0)
    while True:
        timestamp_text = timestamp.strftime("%Y-%m-%d_%H-%M-%S")
        request_path = REQUESTS_DIR / f"Заявка_СИЗ_{timestamp_text}.xlsx"
        if not request_path.exists():
            return request_path
        timestamp += timedelta(seconds=1)


def create_request_file(request_data, request_time=None):
    missing_fields = validate_request_data(request_data)
    if missing_fields:
        raise ValueError(f"Заполните обязательное поле: {missing_fields[0]}")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заявка"
    sheet.merge_cells("A1:B1")
    title_cell = sheet["A1"]
    title_cell.value = REQUEST_TITLE
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    title_cell.fill = PatternFill("solid", fgColor="D9EAF7")
    sheet.row_dimensions[1].height = 36

    thin_side = Side(style="thin", color="808080")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for row_number, field in enumerate(REQUEST_FIELDS, start=3):
        label_cell = sheet.cell(row=row_number, column=1, value=field)
        value = "" if field in REQUEST_REVIEW_FIELDS else request_data.get(field, "")
        value_cell = sheet.cell(row=row_number, column=2, value=str(value or "").strip())

        label_cell.font = Font(bold=True)
        label_cell.fill = PatternFill("solid", fgColor="EAF2F8")
        label_cell.alignment = Alignment(vertical="top", wrap_text=True)
        value_cell.alignment = Alignment(vertical="top", wrap_text=True)
        label_cell.border = border
        value_cell.border = border

        if field in REQUEST_MULTILINE_FIELDS or field == "Комментарий специалиста по ОТ":
            sheet.row_dimensions[row_number].height = 55
        else:
            sheet.row_dimensions[row_number].height = 28

    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 72
    sheet.freeze_panes = "A3"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:B{len(REQUEST_FIELDS) + 2}"

    REQUESTS_DIR.mkdir(parents=True, exist_ok=True)
    request_path = get_request_path(request_time)
    workbook.save(request_path)
    workbook.close()
    return request_path


def ensure_working_directories():
    """Create writable portable folders without creating required resources."""
    for directory in (OUTPUT_DIR, BACKUPS_DIR, REQUESTS_DIR):
        directory.mkdir(parents=True, exist_ok=True)


class SizCardApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Генератор карточек СИЗ")
        self.root.minsize(720, 570)
        self.root.resizable(True, True)

        self.department_var = tk.StringVar()
        self.position_var = tk.StringVar()
        self.field_vars = {label: tk.StringVar() for label in FIELD_LABELS}
        self.departments = []
        self.positions = []
        self.files_ready = False
        self.last_created_card_path = None
        self.print_thread = None
        self.editable_entries = []

        self.configure_styles()
        self.build_form()
        self.center_window(760, 600)
        self.files_ready = self.check_required_files()
        if self.files_ready:
            self.load_departments()
        else:
            self.create_button["state"] = "disabled"

    def configure_styles(self):
        self.style = ttk.Style(self.root)
        if "clam" in self.style.theme_names():
            self.style.theme_use("clam")

        app_background = "#F3F5F7"
        surface_background = "#FFFFFF"
        text_color = "#2B3035"
        muted_color = "#68727D"
        accent_color = "#2F6FAE"
        accent_active = "#285F95"
        neutral_background = "#EDF1F4"
        neutral_active = "#E1E7EC"
        neutral_pressed = "#D4DCE3"
        neutral_border = "#B8C2CB"

        self.root.configure(background=app_background)
        self.style.configure(".", font=("Segoe UI", 10))
        self.style.configure("App.TFrame", background=app_background)
        self.style.configure("Surface.TFrame", background=surface_background)
        self.style.configure(
            "Header.TLabel",
            background=app_background,
            foreground=text_color,
            font=("Segoe UI", 16, "bold"),
        )
        self.style.configure(
            "Subtitle.TLabel",
            background=app_background,
            foreground=muted_color,
            font=("Segoe UI", 10),
        )
        self.style.configure(
            "Section.TLabel",
            background=surface_background,
            foreground=text_color,
            font=("Segoe UI", 11, "bold"),
        )
        self.style.configure(
            "Field.TLabel",
            background=surface_background,
            foreground=text_color,
            font=("Segoe UI", 10),
        )
        self.style.configure(
            "Author.TLabel",
            background=app_background,
            foreground="#7B858E",
            font=("Segoe UI", 8),
            borderwidth=0,
            relief="flat",
        )
        self.style.configure(
            "Neutral.TButton",
            background=neutral_background,
            foreground=text_color,
            font=("Segoe UI", 10),
            padding=(10, 7),
            borderwidth=1,
            bordercolor=neutral_border,
            lightcolor=neutral_border,
            darkcolor=neutral_border,
            relief="raised",
        )
        self.style.map(
            "Neutral.TButton",
            background=[
                ("disabled", "#F2F4F6"),
                ("pressed", neutral_pressed),
                ("active", neutral_active),
            ],
            foreground=[("disabled", "#858F99"), ("!disabled", text_color)],
            bordercolor=[
                ("disabled", "#D3D9DE"),
                ("pressed", "#98A6B2"),
                ("active", "#A8B4BE"),
                ("!disabled", neutral_border),
            ],
            relief=[("pressed", "sunken"), ("!pressed", "raised")],
        )
        self.style.configure(
            "Accent.TButton",
            background=accent_color,
            foreground="#FFFFFF",
            font=("Segoe UI", 10, "bold"),
            padding=(14, 8),
            borderwidth=1,
        )
        self.style.map(
            "Accent.TButton",
            background=[
                ("disabled", "#AAB7C4"),
                ("pressed", accent_active),
                ("active", accent_active),
            ],
            foreground=[("disabled", "#EEF2F5"), ("!disabled", "#FFFFFF")],
        )

    def center_window(self, width, height):
        self.root.update_idletasks()
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        x = max((screen_width - width) // 2, 0)
        y = max((screen_height - height) // 2, 0)
        self.root.geometry(f"{width}x{height}+{x}+{y}")

    def build_form(self):
        self.root.rowconfigure(0, weight=1)
        self.root.columnconfigure(0, weight=1)

        frame = ttk.Frame(self.root, style="App.TFrame", padding=(20, 8))
        frame.grid(row=0, column=0, sticky="nsew")
        frame.columnconfigure(0, weight=1)

        header_frame = ttk.Frame(frame, style="App.TFrame")
        header_frame.grid(row=0, column=0, sticky="ew")
        header_frame.columnconfigure(0, weight=1)
        ttk.Label(
            header_frame,
            text="Генератор карточек СИЗ",
            style="Header.TLabel",
        ).grid(row=0, column=0, sticky="w")
        ttk.Label(
            header_frame,
            text="Создание личной карточки учёта выдачи средств индивидуальной защиты",
            style="Subtitle.TLabel",
        ).grid(row=1, column=0, sticky="w", pady=(2, 0))

        ttk.Separator(frame, orient="horizontal").grid(
            row=1,
            column=0,
            sticky="ew",
            pady=(8, 8),
        )

        employee_frame = ttk.Frame(
            frame,
            style="Surface.TFrame",
            padding=(16, 8),
        )
        employee_frame.grid(row=2, column=0, sticky="ew")
        employee_frame.columnconfigure(0, minsize=150)
        employee_frame.columnconfigure(1, weight=1)
        ttk.Label(
            employee_frame,
            text="Данные работника",
            style="Section.TLabel",
        ).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 5))

        ttk.Label(
            employee_frame,
            text="Подразделение",
            style="Field.TLabel",
        ).grid(row=1, column=0, sticky="w", padx=(0, 14), pady=2)
        self.department_combo = ttk.Combobox(
            employee_frame,
            textvariable=self.department_var,
        )
        self.department_combo.grid(row=1, column=1, sticky="ew", pady=2)
        self.department_combo.bind("<<ComboboxSelected>>", self.on_department_selected)
        self.department_combo.bind("<KeyRelease>", self.on_department_typed)

        ttk.Label(
            employee_frame,
            text="Должность",
            style="Field.TLabel",
        ).grid(row=2, column=0, sticky="w", padx=(0, 14), pady=2)
        self.position_combo = ttk.Combobox(
            employee_frame,
            textvariable=self.position_var,
        )
        self.position_combo.grid(row=2, column=1, sticky="ew", pady=2)
        self.position_combo.bind("<KeyRelease>", self.on_position_typed)

        employee_labels = FIELD_LABELS[:4]
        for index, label in enumerate(employee_labels, start=3):
            display_label = {
                "Номер карточки": "Табельный номер",
                "Дата приема": "Дата приёма",
            }.get(label, label)
            ttk.Label(
                employee_frame,
                text=display_label,
                style="Field.TLabel",
            ).grid(row=index, column=0, sticky="w", padx=(0, 14), pady=2)
            if label == "Пол":
                field = ttk.Combobox(
                    employee_frame,
                    textvariable=self.field_vars[label],
                    values=["М", "Ж"],
                    state="readonly",
                )
            else:
                field = ttk.Entry(employee_frame, textvariable=self.field_vars[label])
            field.grid(row=index, column=1, sticky="ew", pady=2)
            if isinstance(field, ttk.Entry):
                self.configure_editable_entry(field)
            if label == "Дата приема":
                self.date_entry = field
                field.bind("<KeyRelease>", self.format_date_entry, add="+")
                field.bind("<<Paste>>", self.schedule_date_format, add="+")

        size_frame = ttk.Frame(
            frame,
            style="Surface.TFrame",
            padding=(16, 8),
        )
        size_frame.grid(row=3, column=0, sticky="ew", pady=(10, 0))
        size_frame.columnconfigure(0, weight=1, uniform="size_columns")
        size_frame.columnconfigure(1, minsize=24)
        size_frame.columnconfigure(2, weight=1, uniform="size_columns")
        ttk.Label(
            size_frame,
            text="Размерные данные",
            style="Section.TLabel",
        ).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 5))

        left_size_frame = ttk.Frame(size_frame, style="Surface.TFrame")
        left_size_frame.grid(row=1, column=0, sticky="nsew")
        right_size_frame = ttk.Frame(size_frame, style="Surface.TFrame")
        right_size_frame.grid(row=1, column=2, sticky="nsew")
        for column_frame in (left_size_frame, right_size_frame):
            column_frame.columnconfigure(0, minsize=185)
            column_frame.columnconfigure(1, weight=1)

        left_size_fields = ["Рост", "Размер одежды (верх)", "Размер одежды (низ)"]
        right_size_fields = ["Размер обуви", "Размер головного убора", "Размер перчаток"]
        for row, label in enumerate(left_size_fields):
            ttk.Label(
                left_size_frame,
                text=label,
                style="Field.TLabel",
            ).grid(row=row, column=0, sticky="w", padx=(0, 14), pady=2)
            field = ttk.Entry(
                left_size_frame,
                textvariable=self.field_vars[label],
            )
            field.grid(row=row, column=1, sticky="ew", pady=2)
            self.configure_editable_entry(field)

        for row, label in enumerate(right_size_fields):
            ttk.Label(
                right_size_frame,
                text=label,
                style="Field.TLabel",
            ).grid(row=row, column=0, sticky="w", padx=(0, 14), pady=2)
            field = ttk.Entry(
                right_size_frame,
                textvariable=self.field_vars[label],
            )
            field.grid(row=row, column=1, sticky="ew", pady=2)
            self.configure_editable_entry(field)

        action_frame = ttk.Frame(frame, style="App.TFrame")
        action_frame.grid(row=4, column=0, sticky="ew", pady=(10, 0))
        action_frame.columnconfigure(1, weight=1)

        ttk.Button(
            action_frame,
            text="Очистить форму",
            command=self.clear_form,
            style="Neutral.TButton",
        ).grid(
            row=0,
            column=0,
            sticky="w",
        )
        right_actions = ttk.Frame(action_frame, style="App.TFrame")
        right_actions.grid(row=0, column=1, sticky="e")
        self.open_card_button = ttk.Button(
            right_actions,
            text="Открыть папку карточек",
            command=self.open_cards_folder,
            style="Neutral.TButton",
        )
        self.open_card_button.grid(row=0, column=1, padx=(0, 8))
        self.print_card_button = ttk.Button(
            right_actions,
            text="Печать карточки СИЗ",
            command=self.print_created_card,
            state="disabled",
            style="Neutral.TButton",
        )
        self.print_card_button.grid(row=0, column=2, padx=(0, 8))
        self.create_button = ttk.Button(
            right_actions,
            text="Создать карточку",
            command=self.create_card,
            style="Accent.TButton",
        )
        self.create_button.grid(row=0, column=3)

        request_frame = ttk.Frame(
            frame,
            style="Surface.TFrame",
            padding=(16, 8),
        )
        request_frame.grid(row=6, column=0, sticky="ew", pady=(18, 0))
        request_frame.columnconfigure(0, weight=1, uniform="request_buttons")
        request_frame.columnconfigure(1, weight=1, uniform="request_buttons")
        ttk.Label(
            request_frame,
            text="Заявки на изменение норм",
            style="Section.TLabel",
        ).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 8))
        ttk.Button(
            request_frame,
            text="Открыть папку заявок",
            command=self.open_requests_folder,
            style="Neutral.TButton",
        ).grid(row=1, column=0, sticky="ew", padx=(0, 5))
        ttk.Button(
            request_frame,
            text="Заявка на добавление должности",
            command=self.open_request_form,
            style="Neutral.TButton",
        ).grid(row=1, column=1, sticky="ew", padx=(5, 0))
        ttk.Label(
            frame,
            text="Mikhail Kochergin",
            style="Author.TLabel",
        ).grid(row=7, column=0, sticky="e", pady=(12, 2))

    def configure_editable_entry(self, entry):
        self.editable_entries.append(entry)
        entry.bind("<Control-KeyPress>", self.handle_control_shortcut, add="+")
        entry.bind("<Control-a>", lambda event: self.run_entry_shortcut(event, "<<SelectAll>>"), add="+")
        entry.bind("<Control-c>", lambda event: self.run_entry_shortcut(event, "<<Copy>>"), add="+")
        entry.bind("<Control-v>", lambda event: self.run_entry_shortcut(event, "<<Paste>>"), add="+")
        entry.bind("<Control-x>", lambda event: self.run_entry_shortcut(event, "<<Cut>>"), add="+")

    def run_entry_shortcut(self, event, virtual_event):
        if virtual_event == "<<SelectAll>>":
            event.widget.selection_range(0, tk.END)
            event.widget.icursor(tk.END)
        else:
            event.widget.event_generate(virtual_event)
        return "break"

    def handle_control_shortcut(self, event):
        shortcut_by_keycode = {
            65: "<<SelectAll>>",
            67: "<<Copy>>",
            86: "<<Paste>>",
            88: "<<Cut>>",
        }
        shortcut_by_keysym = {
            "a": "<<SelectAll>>",
            "c": "<<Copy>>",
            "v": "<<Paste>>",
            "x": "<<Cut>>",
        }
        virtual_event = shortcut_by_keycode.get(event.keycode)
        if not virtual_event:
            virtual_event = shortcut_by_keysym.get(str(event.keysym).lower())
        if not virtual_event:
            return None

        return self.run_entry_shortcut(event, virtual_event)

    def schedule_date_format(self, _event=None):
        self.date_entry.after_idle(self.format_date_entry)

    def format_date_entry(self, _event=None):
        raw_value = self.date_entry.get()
        cursor_position = self.date_entry.index(tk.INSERT)
        digits_before_cursor = sum(character.isdigit() for character in raw_value[:cursor_position])
        digits = "".join(character for character in raw_value if character.isdigit())[:8]

        if len(digits) <= 2:
            formatted_value = digits
        elif len(digits) <= 4:
            formatted_value = f"{digits[:2]}.{digits[2:]}"
        else:
            formatted_value = f"{digits[:2]}.{digits[2:4]}.{digits[4:]}"

        if formatted_value == raw_value:
            return

        self.field_vars["Дата приема"].set(formatted_value)
        if digits_before_cursor == 0:
            new_cursor_position = 0
        else:
            seen_digits = 0
            new_cursor_position = len(formatted_value)
            for index, character in enumerate(formatted_value, start=1):
                if character.isdigit():
                    seen_digits += 1
                if seen_digits == digits_before_cursor:
                    new_cursor_position = index
                    break
        self.date_entry.icursor(new_cursor_position)

    def check_required_files(self):
        missing_files = []
        if not NORMS_PATH.exists():
            missing_files.append(f"Файл норм СИЗ: {NORMS_PATH}")
        if not TEMPLATE_PATH.exists():
            missing_files.append(f"Шаблон карточки: {TEMPLATE_PATH}")

        if missing_files:
            messagebox.showerror(
                "Не найдены обязательные файлы",
                "Программа не может создать карточку, потому что не найдены обязательные файлы:\n\n"
                + "\n".join(missing_files),
            )
            return False

        return True

    def load_departments(self):
        if not NORMS_PATH.exists():
            messagebox.showerror(
                "Не найден файл норм",
                f"Файл data/norms.xlsx не найден.\nПоложите рабочий файл норм сюда:\n{NORMS_PATH}",
            )
            self.department_combo["values"] = []
            self.position_combo["values"] = []
            return

        try:
            departments = get_departments(NORMS_PATH)
        except Exception as error:
            messagebox.showerror("Ошибка", f"Не удалось прочитать data/norms.xlsx:\n{error}")
            departments = []

        self.departments = departments
        self.department_combo["values"] = departments
        if departments:
            self.department_var.set(departments[0])
            self.update_positions(departments[0], clear_position=True)
        else:
            messagebox.showerror(
                "Нет подразделений",
                "В data/norms.xlsx не найдены подразделения. Проверьте файл норм.",
            )

    def normalize_choice(self, value):
        return str(value or "").strip().casefold()

    def filter_values(self, values, search_text):
        search = self.normalize_choice(search_text)
        if not search:
            return values
        return [value for value in values if search in self.normalize_choice(value)]

    def find_exact_value(self, values, search_text):
        search = str(search_text or "").strip()
        for value in values:
            if str(value).strip() == search:
                return value
        return None

    def update_combo_values(self, combo, values, show_list=False):
        combo["values"] = values
        if show_list and values:
            try:
                combo.after_idle(lambda: combo.event_generate("<Down>"))
            except AttributeError:
                pass

    def on_department_selected(self, _event=None):
        department = self.find_exact_value(self.departments, self.department_var.get())
        if department:
            self.department_var.set(department)
            self.update_positions(department, clear_position=True)

    def on_department_typed(self, event=None):
        if event and event.keysym in NAVIGATION_KEYS:
            return

        text = self.department_var.get()
        matches = self.filter_values(self.departments, text)
        self.update_combo_values(self.department_combo, matches, show_list=True)

        department = self.find_exact_value(self.departments, text)
        if department:
            self.update_positions(department, clear_position=False, show_empty_message=False)
        else:
            self.positions = []
            self.update_combo_values(self.position_combo, [])
            self.position_var.set("")

    def on_position_typed(self, event=None):
        if event and event.keysym in NAVIGATION_KEYS:
            return

        department = self.find_exact_value(self.departments, self.department_var.get())
        if not department:
            self.update_combo_values(self.position_combo, [])
            return

        if not self.positions:
            self.update_positions(department, clear_position=False, show_empty_message=False)
        matches = self.filter_values(self.positions, self.position_var.get())
        self.update_combo_values(self.position_combo, matches, show_list=True)

    def update_positions(self, department, clear_position=True, show_empty_message=True):
        try:
            positions = get_positions_by_department(NORMS_PATH, department)
        except Exception as error:
            messagebox.showerror("Ошибка", f"Не удалось прочитать должности:\n{error}")
            positions = []

        self.positions = positions
        self.update_combo_values(self.position_combo, positions)

        current_position = self.find_exact_value(positions, self.position_var.get())
        if clear_position:
            self.position_var.set("")
        elif current_position:
            self.position_var.set(current_position)
        else:
            self.position_var.set("")

        if department and not positions and show_empty_message:
            messagebox.showerror(
                "Нет должностей",
                f"В data/norms.xlsx нет должностей для подразделения:\n{department}",
            )

    def collect_employee(self):
        employee = {
            "Подразделение": self.department_var.get().strip(),
            "Должность": self.position_var.get().strip(),
        }

        for label, variable in self.field_vars.items():
            if label in {"Размер одежды (верх)", "Размер одежды (низ)"}:
                continue
            employee_column = FIELD_TO_EMPLOYEE_COLUMN.get(label, label)
            employee[employee_column] = variable.get().strip()

        clothing_sizes = [
            self.field_vars["Размер одежды (верх)"].get().strip(),
            self.field_vars["Размер одежды (низ)"].get().strip(),
        ]
        employee["Размер одежды"] = "/".join(size for size in clothing_sizes if size)

        return employee

    def validate_employee(self, employee):
        required_to_employee_column = {
            field: FIELD_TO_EMPLOYEE_COLUMN.get(field, field)
            for field in REQUIRED_FIELDS
        }
        missing = [
            field
            for field, employee_column in required_to_employee_column.items()
            if not employee.get(employee_column)
        ]
        if missing:
            messagebox.showerror("Заполните поле", f"Заполните поле: {missing[0]}")
            return False

        department = self.find_exact_value(self.departments, employee["Подразделение"])
        if not department:
            messagebox.showerror(
                "Проверьте подразделение",
                "Подразделение не найдено в нормах СИЗ. Выберите значение из списка.",
            )
            return False

        positions = get_positions_by_department(NORMS_PATH, department)
        position = self.find_exact_value(positions, employee["Должность"])
        if not position:
            messagebox.showerror(
                "Проверьте должность",
                "Должность не найдена для выбранного подразделения. Выберите значение из списка.",
            )
            return False

        employee["Подразделение"] = department
        employee["Должность"] = position
        return True

    def create_card(self):
        if not self.check_required_files():
            return

        employee = self.collect_employee()
        if not self.validate_employee(employee):
            return

        try:
            output_path = create_single_card(
                employee=employee,
                norms_path=NORMS_PATH,
                template_path=TEMPLATE_PATH,
                output_dir=OUTPUT_DIR,
            )
        except Exception as error:
            messagebox.showerror("Ошибка", f"Не удалось создать карточку. Подробности: {error}")
            return

        if output_path is None:
            messagebox.showerror(
                "Не найдены нормы",
                "Для выбранного подразделения и должности нормы СИЗ не найдены. Проверьте файл data/norms.xlsx.",
            )
            return

        self.last_created_card_path = Path(output_path)
        self.print_card_button["state"] = "normal"
        messagebox.showinfo("Карточка создана", f"Карточка создана:\n{output_path}")

    def open_cards_folder(self):
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        try:
            os.startfile(OUTPUT_DIR)
        except OSError as error:
            messagebox.showerror("Ошибка", f"Не удалось открыть папку карточек:\n{error}")

    def open_requests_folder(self):
        REQUESTS_DIR.mkdir(parents=True, exist_ok=True)
        try:
            os.startfile(REQUESTS_DIR)
        except OSError as error:
            messagebox.showerror("Ошибка", f"Не удалось открыть папку заявок:\n{error}")

    def open_request_form(self):
        window = tk.Toplevel(self.root)
        window.title(REQUEST_TITLE)
        window.resizable(False, False)
        window.transient(self.root)

        frame = ttk.Frame(window, padding=16)
        frame.grid(row=0, column=0, sticky="nsew")

        request_vars = {
            "Дата заявки": tk.StringVar(value=date.today().isoformat()),
            "Подразделение": tk.StringVar(value=self.department_var.get().strip()),
            "Должность": tk.StringVar(value=self.position_var.get().strip()),
            "ФИО инициатора заявки": tk.StringVar(),
            "Контакт инициатора": tk.StringVar(),
        }
        text_fields = {}
        request_state = {"last_created_path": None}
        form_fields = [
            "Подразделение",
            "Должность",
            "Описание выполняемых работ",
            "Основание для добавления в нормы СИЗ",
            "Предполагаемые СИЗ, если известно",
            "ФИО инициатора заявки",
            "Контакт инициатора",
            "Дата заявки",
        ]

        for row_number, field in enumerate(form_fields):
            ttk.Label(frame, text=field).grid(
                row=row_number,
                column=0,
                sticky="nw",
                padx=(0, 12),
                pady=4,
            )
            if field in REQUEST_MULTILINE_FIELDS:
                widget = tk.Text(frame, width=58, height=4, wrap="word")
                widget.grid(row=row_number, column=1, sticky="ew", pady=4)
                text_fields[field] = widget
            else:
                widget = ttk.Entry(frame, textvariable=request_vars[field], width=60)
                widget.grid(row=row_number, column=1, sticky="ew", pady=4)

        def save_request():
            request_data = {
                field: widget.get("1.0", "end").strip()
                for field, widget in text_fields.items()
            }
            request_data.update(
                {
                    field: variable.get().strip()
                    for field, variable in request_vars.items()
                }
            )

            missing_fields = validate_request_data(request_data)
            if missing_fields:
                messagebox.showerror(
                    "Не заполнено обязательное поле",
                    f"Заполните обязательное поле: {missing_fields[0]}",
                    parent=window,
                )
                return

            try:
                request_path = create_request_file(request_data)
            except (OSError, ValueError) as error:
                messagebox.showerror(
                    "Ошибка создания заявки",
                    f"Не удалось создать заявку:\n{error}",
                    parent=window,
                )
                return

            request_state["last_created_path"] = Path(request_path)
            open_request_button["state"] = "normal"
            messagebox.showinfo(
                "Заявка создана",
                f"Заявка создана: {request_path}. "
                "Передайте файл специалисту по охране труда.",
                parent=window,
            )

        def open_created_request():
            request_path = request_state["last_created_path"]
            if not request_path or not request_path.exists():
                messagebox.showerror(
                    "Файл не найден",
                    "Файл заявки не найден. Проверьте папку requests.",
                    parent=window,
                )
                open_request_button["state"] = "disabled"
                return

            try:
                os.startfile(request_path)
            except OSError as error:
                messagebox.showerror(
                    "Ошибка",
                    f"Не удалось открыть заявку:\n{error}",
                    parent=window,
                )

        button_frame = ttk.Frame(frame)
        button_frame.grid(
            row=len(form_fields),
            column=0,
            columnspan=2,
            sticky="e",
            pady=(12, 0),
        )
        ttk.Button(button_frame, text="Закрыть", command=window.destroy).grid(
            row=0,
            column=0,
            padx=(0, 8),
        )
        open_request_button = ttk.Button(
            button_frame,
            text="Открыть заявку",
            command=open_created_request,
            state="disabled",
        )
        open_request_button.grid(
            row=0,
            column=1,
            padx=(0, 8),
        )
        ttk.Button(button_frame, text="Сохранить заявку", command=save_request).grid(
            row=0,
            column=2,
        )

        window.grab_set()
        window.focus_set()

    def clear_form(self):
        for variable in self.field_vars.values():
            variable.set("")
        self.last_created_card_path = None
        self.print_card_button["state"] = "disabled"

    def print_created_card(self):
        if not self.last_created_card_path or not self.last_created_card_path.exists():
            messagebox.showerror(
                "Файл не найден",
                "Созданная карточка не найдена. Возможно, файл был удален или перемещен.",
            )
            self.print_card_button["state"] = "disabled"
            return

        card_path = self.last_created_card_path.resolve()
        self.print_card_button["state"] = "disabled"
        self.print_thread = threading.Thread(
            target=self.run_print_preview,
            args=(card_path,),
            daemon=False,
        )
        self.print_thread.start()

    def run_print_preview(self, card_path):
        environment = os.environ.copy()
        environment["SIZ_CARD_PATH"] = str(card_path)
        encoded_script = base64.b64encode(
            PRINT_PREVIEW_SCRIPT.encode("utf-16le")
        ).decode("ascii")

        try:
            result = subprocess.run(
                [
                    "powershell.exe",
                    "-NoLogo",
                    "-NoProfile",
                    "-NonInteractive",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-EncodedCommand",
                    encoded_script,
                ],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                env=environment,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
                check=False,
            )
        except OSError as error:
            self.schedule_print_preview_result(str(error))
            return

        if result.returncode:
            encoded_error = next(
                (
                    line.removeprefix("ERROR_BASE64=").strip()
                    for line in result.stdout.splitlines()
                    if line.startswith("ERROR_BASE64=")
                ),
                "",
            )
            try:
                detail = base64.b64decode(encoded_error).decode("utf-8")
            except (ValueError, UnicodeDecodeError):
                detail = "Ошибка PowerShell."
            if not detail:
                detail = "Ошибка PowerShell."
            self.schedule_print_preview_result(detail[:500])
            return

        self.schedule_print_preview_result()

    def schedule_print_preview_result(self, error_message=None):
        try:
            self.root.after(0, self.finish_print_preview, error_message)
        except (RuntimeError, tk.TclError):
            pass

    def finish_print_preview(self, error_message=None):
        if self.last_created_card_path and self.last_created_card_path.exists():
            self.print_card_button["state"] = "normal"
        else:
            self.print_card_button["state"] = "disabled"

        if error_message:
            messagebox.showerror(
                "Ошибка печати",
                f"Не удалось открыть карточку для печати.\nПодробности: {error_message}",
            )


def main():
    root = tk.Tk()
    if APP_ICON_PATH.exists():
        try:
            root.iconbitmap(str(APP_ICON_PATH))
        except tk.TclError:
            pass
    try:
        ensure_working_directories()
    except OSError as error:
        messagebox.showerror(
            "Не удалось подготовить рабочие папки",
            "Программа не может создать рабочие папки рядом с приложением:\n\n"
            f"{error}",
            parent=root,
        )
        root.destroy()
        return
    SizCardApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
