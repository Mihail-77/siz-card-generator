from copy import copy
from math import ceil
from pathlib import Path
from re import sub

from openpyxl.cell.cell import MergedCell
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter


NORM_COLUMNS = [
    "Подразделение",
    "Должность",
    "Наименование СИЗ",
    "Норма выдачи",
    "Срок носки",
    "Основание",
    "Примечание",
]

EMPLOYEE_COLUMNS = [
    "Табельный номер",
    "ФИО",
    "Подразделение",
    "Должность",
    "Дата приема",
    "Пол",
    "Рост",
    "Размер одежды",
    "Размер обуви",
    "Размер головного убора",
    "Размер перчаток",
]

FRONT_SHEET_NAME = "Лицевая сторона"
BACK_SHEET_NAME = "Оборотная сторона"
NORM_TABLE_HEADER_ROW = 29
NORM_TABLE_START_ROW = NORM_TABLE_HEADER_ROW + 1
ISSUE_TABLE_HEADER_ROW = 3
ISSUE_TABLE_NUMBER_ROW = 5
ISSUE_TABLE_START_ROW = ISSUE_TABLE_NUMBER_ROW + 1
FRONT_TEMPLATE_ROW = 30
FRONT_TEMPLATE_LAST_ROW = 33
BACK_TEMPLATE_ROW = 6
BACK_TEMPLATE_LAST_ROW = 7
FRONT_TABLE_MERGES = [(1, 48), (49, 66), (67, 87), (88, 108)]
BACK_TABLE_MERGES = [
    (1, 17),
    (18, 35),
    (36, 44),
    (45, 52),
    (53, 60),
    (61, 70),
    (71, 79),
    (80, 87),
    (88, 97),
    (98, 108),
]
BACK_NOTE_TEXTS = [
    "* - информация указывается только для дерматологических СИЗ",
    "** - информация указывается для всех СИЗ, кроме дерматологических СИЗ и СИЗ однократного применения",
]
THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)


def normalize(value):
    """Prepare text for comparison: remove extra spaces and ignore case."""
    if value is None:
        return ""
    return str(value).strip().lower()


def safe_filename(value):
    """Make a simple file name from employee name or personnel number."""
    name = str(value).strip()
    name = sub(r'[\\/:*?"<>|]+', "_", name)
    name = sub(r"\s+", "_", name)
    return name or "employee"


def split_full_name(full_name):
    """Split Russian full name into surname, name and patronymic if possible."""
    parts = str(full_name or "").strip().split()
    surname = parts[0] if len(parts) > 0 else ""
    name = parts[1] if len(parts) > 1 else ""
    patronymic = " ".join(parts[2:]) if len(parts) > 2 else ""
    return surname, name, patronymic


def safe_write_cell(sheet, cell, value):
    """
    Write a value to a cell, respecting merged ranges.

    If the target cell is inside a merged range, openpyxl allows writing only
    to the top-left cell of that range. This helper redirects the write there.
    """
    if isinstance(cell, str):
        target = sheet[cell]
        row = target.row
        column = target.column
    else:
        row, column = cell

    coordinate = f"{get_column_letter(column)}{row}"
    for merged_range in sheet.merged_cells.ranges:
        if coordinate in merged_range:
            sheet.cell(row=merged_range.min_row, column=merged_range.min_col).value = value
            return

    sheet.cell(row=row, column=column).value = value


def find_merged_range(sheet, row, column):
    """Return merged range containing the cell, if it exists."""
    coordinate = f"{get_column_letter(column)}{row}"
    for merged_range in sheet.merged_cells.ranges:
        if coordinate in merged_range:
            return merged_range
    return None


def enable_wrap_text(sheet, row, column):
    """Enable text wrapping without changing the existing alignment."""
    merged_range = find_merged_range(sheet, row, column)
    target_row = merged_range.min_row if merged_range else row
    target_column = merged_range.min_col if merged_range else column
    cell = sheet.cell(row=target_row, column=target_column)

    alignment = copy(cell.alignment)
    alignment.wrap_text = True
    cell.alignment = alignment

    return merged_range


def set_employee_field_row_height(sheet, cell, text):
    """Increase employee info row height for long values in merged fields."""
    target = sheet[cell]
    merged_range = enable_wrap_text(sheet, target.row, target.column)
    column_count = (
        merged_range.max_col - merged_range.min_col + 1
        if merged_range
        else 1
    )
    chars_per_line = max(24, min(42, int(column_count * 0.85)))

    text_value = str(text or "")
    estimated_lines = 0
    for text_part in text_value.splitlines() or [""]:
        estimated_lines += max(1, ceil(len(text_part) / chars_per_line))

    base_height = sheet.row_dimensions[target.row].height or sheet.sheet_format.defaultRowHeight or 15
    if estimated_lines <= 1:
        return

    sheet.row_dimensions[target.row].height = max(base_height, min(estimated_lines * 17, 75))


def copy_row_format(sheet, source_row, target_row, max_column=108):
    """Copy row height and cell formatting from a template row."""
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column in range(1, max_column + 1):
        source_cell = sheet.cell(row=source_row, column=column)
        target_cell = sheet.cell(row=target_row, column=column)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format


def unmerge_ranges_on_row(sheet, row_number):
    """Remove merged ranges that are fully placed on one target row."""
    ranges_to_unmerge = [
        merged_range.coord
        for merged_range in sheet.merged_cells.ranges
        if merged_range.min_row == row_number and merged_range.max_row == row_number
    ]
    for merged_range in ranges_to_unmerge:
        try:
            sheet.unmerge_cells(merged_range)
        except KeyError:
            for existing_range in list(sheet.merged_cells.ranges):
                if existing_range.coord == merged_range:
                    sheet.merged_cells.ranges.remove(existing_range)
                    break


def apply_row_merges(sheet, row_number, column_ranges):
    """Apply horizontal merged ranges for a generated table row."""
    for start_column, end_column in column_ranges:
        sheet.merge_cells(
            start_row=row_number,
            start_column=start_column,
            end_row=row_number,
            end_column=end_column,
        )


def normalize_workbook_font(workbook):
    """Use Times New Roman 11 as the default font in generated workbooks."""
    for named_style in workbook._named_styles:
        if named_style.name == "Normal":
            named_style.font = Font(name="Times New Roman", size=11)
            break


def estimate_text_height(text, chars_per_line, base_height=24, line_height=17):
    """Estimate row height for wrapped text."""
    visual_lines = 0
    for text_part in str(text or "").splitlines() or [""]:
        visual_lines += max(1, ceil(len(text_part) / chars_per_line))
    if visual_lines <= 1:
        return base_height
    return max(base_height, visual_lines * line_height)


def set_table_row_height(sheet, row_number, row_values):
    """Increase prescribed PPE table row height by the longest visible field."""
    base_height = sheet.row_dimensions[row_number].height or 24
    fields = [
        (1, row_values.get("Наименование СИЗ"), 42),
        (49, row_values.get("Основание"), 20),
        (67, row_values.get("Срок носки"), 22),
        (88, row_values.get("Норма выдачи"), 14),
    ]

    row_height = base_height
    for column, text, chars_per_line in fields:
        cell = sheet.cell(row=row_number, column=column)
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        cell.alignment = alignment
        cell.font = Font(name="Times New Roman", size=11)
        row_height = max(row_height, estimate_text_height(text, chars_per_line, base_height))

    sheet.row_dimensions[row_number].height = row_height


def set_back_table_row_height(sheet, row_number, text):
    """Set row height for the issue/return table based on the PPE name length."""
    text_value = str(text or "")
    text_length = len(text_value)
    visual_lines = 0
    for text_part in text_value.splitlines() or [""]:
        visual_lines += max(1, ceil(len(text_part) / 14))

    row_height = max(36, visual_lines * 17)
    if 25 <= text_length <= 40:
        row_height = max(row_height, 54)

    sheet.row_dimensions[row_number].height = min(row_height, 110)
    sheet.cell(row=row_number, column=1).alignment = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )


def style_back_table_row(sheet, row_number):
    """Apply table style to every cell in one issue/return row."""
    for column in range(1, 109):
        cell = sheet.cell(row=row_number, column=column)
        cell.font = Font(name="Times New Roman", size=11)
        cell.border = copy(THIN_BORDER)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for column in range(1, 18):
        sheet.cell(row=row_number, column=column).alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )


def prepare_front_table_rows(sheet, row_count):
    """Prepare formatted rows for the prescribed PPE table."""
    if row_count <= 0:
        return

    existing_rows = FRONT_TEMPLATE_LAST_ROW - FRONT_TEMPLATE_ROW + 1
    extra_rows = max(0, row_count - existing_rows)
    if extra_rows:
        sheet.insert_rows(FRONT_TEMPLATE_LAST_ROW + 1, extra_rows)

    for row_number in range(FRONT_TEMPLATE_ROW, FRONT_TEMPLATE_ROW + row_count):
        unmerge_ranges_on_row(sheet, row_number)
        remove_stale_merged_cells(sheet, row_number)
        copy_row_format(sheet, FRONT_TEMPLATE_ROW, row_number)
        apply_row_merges(sheet, row_number, FRONT_TABLE_MERGES)


def prepare_back_table_rows(sheet, row_count):
    """Prepare formatted rows for the PPE issue and return table."""
    if row_count <= 0:
        return

    existing_rows = BACK_TEMPLATE_LAST_ROW - BACK_TEMPLATE_ROW + 1
    extra_rows = max(0, row_count - existing_rows)
    if extra_rows:
        sheet.insert_rows(BACK_TEMPLATE_LAST_ROW + 1, extra_rows)

    for row_number in range(BACK_TEMPLATE_ROW, BACK_TEMPLATE_ROW + row_count):
        unmerge_ranges_on_row(sheet, row_number)
        remove_stale_merged_cells(sheet, row_number)
        copy_row_format(sheet, BACK_TEMPLATE_ROW, row_number)
        style_back_table_row(sheet, row_number)
        apply_row_merges(sheet, row_number, BACK_TABLE_MERGES)


def clear_merged_ranges_intersecting_rows(sheet, start_row, end_row):
    """Remove merged ranges that intersect the target rows."""
    ranges_to_unmerge = [
        merged_range.coord
        for merged_range in sheet.merged_cells.ranges
        if merged_range.min_row <= end_row and merged_range.max_row >= start_row
    ]
    for merged_range in ranges_to_unmerge:
        sheet.unmerge_cells(merged_range)


def remove_stale_merged_cells(sheet, row_number, max_column=108):
    """Remove leftover MergedCell objects after row insertion/unmerge."""
    for column in range(1, max_column + 1):
        key = (row_number, column)
        if isinstance(sheet._cells.get(key), MergedCell):
            del sheet._cells[key]


def prepare_back_notes(sheet, norms_count):
    """Place note rows directly below the generated issue/return table."""
    first_note_row = ISSUE_TABLE_START_ROW + norms_count
    second_note_row = first_note_row + 1

    clear_merged_ranges_intersecting_rows(sheet, first_note_row, second_note_row)

    for row_number, note_text in zip([first_note_row, second_note_row], BACK_NOTE_TEXTS):
        remove_stale_merged_cells(sheet, row_number)
        for column in range(1, 109):
            cell = sheet.cell(row=row_number, column=column)
            cell.value = None
            cell.border = copy(sheet["A1"].border)
            cell.fill = copy(sheet["A1"].fill)

        sheet.merge_cells(start_row=row_number, start_column=2, end_row=row_number, end_column=108)
        note_cell = sheet.cell(row=row_number, column=2)
        note_cell.value = note_text
        note_cell.font = Font(name="Times New Roman", size=11)
        note_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.row_dimensions[row_number].height = 18 if row_number == first_note_row else 34


def update_print_areas(front_sheet, back_sheet):
    """Keep print areas aligned with rows inserted into generated cards."""
    front_sheet.print_area = f"A1:DD{front_sheet.max_row}"
    back_sheet.print_area = f"A1:DD{back_sheet.max_row}"


def read_table(path, required_columns):
    """Read an Excel sheet where the first row contains column headers."""
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active

    headers = [cell.value for cell in sheet[1]]
    missing_columns = [column for column in required_columns if column not in headers]
    if missing_columns:
        joined = ", ".join(missing_columns)
        raise ValueError(f"В файле {path.name} нет обязательных колонок: {joined}")

    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in row):
            continue

        item = {}
        for column in required_columns:
            index = headers.index(column)
            item[column] = row[index]
        rows.append(item)

    return rows


def load_norms(norms_path):
    """Group norms by department and position."""
    rows = read_table(norms_path, NORM_COLUMNS)
    norms_by_key = {}

    for row in rows:
        department = normalize(row["Подразделение"])
        position = normalize(row["Должность"])
        key = (department, position)
        norms_by_key.setdefault(key, []).append(row)

    return norms_by_key


def get_departments(norms_path):
    """Return unique departments from norms.xlsx."""
    rows = read_table(norms_path, NORM_COLUMNS)
    departments = {
        str(row["Подразделение"]).strip()
        for row in rows
        if row["Подразделение"] is not None and str(row["Подразделение"]).strip()
    }
    return sorted(departments)


def get_positions_by_department(norms_path, department):
    """Return positions available for the selected department."""
    rows = read_table(norms_path, NORM_COLUMNS)
    selected_department = normalize(department)
    positions = {
        str(row["Должность"]).strip()
        for row in rows
        if normalize(row["Подразделение"]) == selected_department
        and row["Должность"] is not None
        and str(row["Должность"]).strip()
    }
    return sorted(positions)


def find_norms_for_employee(norms_by_key, employee):
    """Find PPE norms by department and position."""
    department = normalize(employee["Подразделение"])
    position = normalize(employee["Должность"])
    return norms_by_key.get((department, position))


def build_warning(employee, reason):
    """Create one warning row for warnings.xlsx."""
    return {
        "Табельный номер": employee["Табельный номер"],
        "ФИО": employee["ФИО"],
        "Подразделение": employee["Подразделение"],
        "Должность": employee["Должность"],
        "Причина": reason,
    }


def style_generated_card(front_sheet, back_sheet, norms_count):
    """Apply basic formatting to the generated card."""
    last_norm_row = max(NORM_TABLE_HEADER_ROW, NORM_TABLE_START_ROW + norms_count - 1)
    last_issue_row = max(ISSUE_TABLE_START_ROW, ISSUE_TABLE_START_ROW + norms_count - 1)

    for row in front_sheet.iter_rows(min_row=NORM_TABLE_START_ROW, max_row=last_norm_row, min_col=1, max_col=108):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row in back_sheet.iter_rows(min_row=ISSUE_TABLE_START_ROW, max_row=last_issue_row, min_col=1, max_col=108):
        for cell in row:
            horizontal = "left" if cell.column <= 17 else "center"
            cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)


def create_card(employee, norms, template_path, output_dir):
    """Create one personal PPE card from the template."""
    workbook = load_workbook(template_path)
    normalize_workbook_font(workbook)
    front_sheet = workbook[FRONT_SHEET_NAME] if FRONT_SHEET_NAME in workbook.sheetnames else workbook.active
    back_sheet = workbook[BACK_SHEET_NAME] if BACK_SHEET_NAME in workbook.sheetnames else workbook.create_sheet(BACK_SHEET_NAME)

    full_name = employee["ФИО"]
    personnel_number = employee["Табельный номер"]
    surname, name, patronymic = split_full_name(full_name)

    safe_write_cell(front_sheet, "AF16", f"ЛИЧНАЯ КАРТОЧКА № {personnel_number}")
    safe_write_cell(front_sheet, "L19", surname)
    safe_write_cell(front_sheet, "G20", name)
    safe_write_cell(front_sheet, "AV20", patronymic)
    safe_write_cell(front_sheet, "T21", personnel_number)
    safe_write_cell(front_sheet, "AE22", employee["Подразделение"])
    safe_write_cell(front_sheet, "AA23", employee["Должность"])
    set_employee_field_row_height(front_sheet, "AE22", employee["Подразделение"])
    set_employee_field_row_height(front_sheet, "AA23", employee["Должность"])
    safe_write_cell(front_sheet, "AE24", employee["Дата приема"])
    safe_write_cell(front_sheet, "BW19", employee["Пол"])
    safe_write_cell(front_sheet, "BX20", employee["Рост"])
    safe_write_cell(front_sheet, "CA22", employee["Размер одежды"])
    safe_write_cell(front_sheet, "BY23", employee["Размер обуви"])
    safe_write_cell(front_sheet, "CJ24", employee["Размер головного убора"])
    safe_write_cell(front_sheet, "CB27", employee["Размер перчаток"])

    prepare_front_table_rows(front_sheet, len(norms))
    prepare_back_table_rows(back_sheet, len(norms))

    for index, norm in enumerate(norms, start=1):
        norm_row = NORM_TABLE_START_ROW + index - 1
        issue_row = ISSUE_TABLE_START_ROW + index - 1

        safe_write_cell(front_sheet, f"A{norm_row}", norm["Наименование СИЗ"])
        safe_write_cell(front_sheet, f"AW{norm_row}", norm["Основание"])
        safe_write_cell(front_sheet, f"BO{norm_row}", norm["Срок носки"])
        safe_write_cell(front_sheet, f"CJ{norm_row}", norm["Норма выдачи"])
        set_table_row_height(front_sheet, norm_row, norm)

        safe_write_cell(back_sheet, f"A{issue_row}", norm["Наименование СИЗ"])
        for column in range(18, 109):
            safe_write_cell(back_sheet, (issue_row, column), None)
        set_back_table_row_height(back_sheet, issue_row, norm["Наименование СИЗ"])

    style_generated_card(front_sheet, back_sheet, len(norms))
    prepare_back_notes(back_sheet, len(norms))
    update_print_areas(front_sheet, back_sheet)

    output_dir.mkdir(parents=True, exist_ok=True)
    file_name = f"{safe_filename(personnel_number)}_{safe_filename(full_name)}.xlsx"
    workbook.save(output_dir / file_name)


def create_single_card(employee, norms_path, template_path, output_dir):
    """
    Create one card from employee data passed from code.

    This function is intended for a future GUI where employee data will be
    entered on the screen instead of being read from employees.xlsx.
    """
    missing_columns = [column for column in EMPLOYEE_COLUMNS if column not in employee]
    if missing_columns:
        joined = ", ".join(missing_columns)
        raise ValueError(f"Не хватает данных работника: {joined}")

    norms_by_key = load_norms(norms_path)
    norms = find_norms_for_employee(norms_by_key, employee)
    if not norms:
        reason = "Не найдены нормы для подразделения и должности"
        save_warnings([build_warning(employee, reason)], output_dir)
        return None

    create_card(employee, norms, template_path, output_dir)
    full_name = employee["ФИО"]
    personnel_number = employee["Табельный номер"]
    file_name = f"{safe_filename(personnel_number)}_{safe_filename(full_name)}.xlsx"
    return output_dir / file_name


def save_warnings(warnings, output_dir):
    """Save warnings to output/warnings.xlsx."""
    output_dir.mkdir(parents=True, exist_ok=True)
    warnings_path = output_dir / "warnings.xlsx"

    if not warnings:
        if warnings_path.exists():
            warnings_path.unlink()
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Предупреждения"
    sheet.append(["Табельный номер", "ФИО", "Подразделение", "Должность", "Причина"])

    for warning in warnings:
        sheet.append(
            [
                warning["Табельный номер"],
                warning["ФИО"],
                warning["Подразделение"],
                warning["Должность"],
                warning["Причина"],
            ]
        )

    for column in ["A", "B", "C", "D", "E"]:
        sheet.column_dimensions[column].width = 24

    workbook.save(warnings_path)


def run(norms_path, employees_path, template_path, output_dir):
    """Run the whole generation process."""
    norms_by_key = load_norms(norms_path)
    employees = read_table(employees_path, EMPLOYEE_COLUMNS)

    warnings = []
    created_count = 0

    for employee in employees:
        norms = find_norms_for_employee(norms_by_key, employee)

        if not norms:
            reason = "Не найдены нормы для подразделения и должности"
            warnings.append(build_warning(employee, reason))
            continue

        create_card(employee, norms, template_path, output_dir)
        created_count += 1

    save_warnings(warnings, output_dir)

    print(f"Готово. Создано карточек: {created_count}")
    print(f"Предупреждений: {len(warnings)}")
    print(f"Папка результата: {output_dir}")
