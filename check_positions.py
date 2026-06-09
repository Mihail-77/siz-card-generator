from difflib import SequenceMatcher
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "output"
POSITIONS_PATH = DATA_DIR / "positions.xlsx"
NORMS_PATH = DATA_DIR / "norms.xlsx"
REPORT_PATH = OUTPUT_DIR / "positions_check.xlsx"
POSSIBLE_MATCH_COMMENT = (
    "Возможное переименование или различие в написании. Требуется проверка."
)
GRADE_MATCH_COMMENT = (
    "Должность из штатки входит в групповую норму по разряду/категории."
)
GRADE_PATTERN = re.compile(
    r"^(?P<base>.*?)"
    r"(?P<numbers>\d+(?:\s*,\s*\d+)*)"
    r"(?:\s*,)?\s+"
    r"(?P<type>разряд(?:а|ы)?|категори(?:я|и))$",
    re.IGNORECASE,
)

POSITIONS_COLUMNS = [
    "Подразделение",
    "Должность",
    "СИЗ предусмотрены",
    "Комментарий",
]
NORM_COLUMNS = [
    "Подразделение",
    "Должность",
    "Наименование СИЗ",
    "Норма выдачи",
    "Срок носки",
    "Основание",
    "Примечание",
]


def normalize(value):
    """Prepare text for comparison without changing source values."""
    return str(value or "").strip().casefold()


def pair_key(row):
    return (normalize(row["Подразделение"]), normalize(row["Должность"]))


def style_header(sheet):
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(sheet):
    widths = {
        "A": 38,
        "B": 48,
        "C": 20,
        "D": 48,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def set_possible_matches_widths(sheet):
    widths = {
        "A": 38,
        "B": 48,
        "C": 38,
        "D": 48,
        "E": 24,
        "F": 22,
        "G": 68,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def set_grade_matches_widths(sheet):
    widths = {
        "A": 38,
        "B": 52,
        "C": 58,
        "D": 16,
        "E": 20,
        "F": 24,
        "G": 68,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def create_positions_placeholder(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Должности"
    sheet.append(POSITIONS_COLUMNS)
    sheet.append(
        [
            "Административно-хозяйственный отдел (АХО)",
            "Начальник отдела",
            "Да",
            "Учебная строка для проверки сверки, заменить на штатное расписание",
        ]
    )
    sheet.append(
        [
            "Транспортный отдел (ТО)",
            "Водитель-экспедитор",
            "Да",
            "Учебная строка для проверки сверки, заменить на штатное расписание",
        ]
    )
    sheet.append(
        [
            "Учебное подразделение",
            "Учебная должность без норм СИЗ",
            "Нет",
            "Демонстрационная строка, показывает раздел отчета 'нет в нормах'",
        ]
    )

    style_header(sheet)
    set_widths(sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def read_table(path, required_columns):
    if not path.exists():
        raise FileNotFoundError(f"Не найден файл: {path}")

    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    missing_columns = [column for column in required_columns if column not in headers]
    if missing_columns:
        joined = ", ".join(missing_columns)
        raise ValueError(f"В файле {path} нет обязательных колонок: {joined}")

    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in row):
            continue

        item = {}
        for column in required_columns:
            item[column] = row[headers.index(column)]
        rows.append(item)

    return rows


def unique_pairs(rows, include_extra=False):
    pairs = {}
    for row in rows:
        key = pair_key(row)
        if not key[0] or not key[1]:
            continue
        if key not in pairs:
            pairs[key] = row if include_extra else {
                "Подразделение": row["Подразделение"],
                "Должность": row["Должность"],
            }
    return pairs


def append_report_rows(sheet, rows):
    for row in rows:
        sheet.append(
            [
                row.get("Подразделение", ""),
                row.get("Должность", ""),
                row.get("СИЗ предусмотрены", ""),
                row.get("Комментарий", ""),
            ]
        )


def similarity_percent(left, right):
    """Return case-insensitive text similarity as a percentage."""
    return round(SequenceMatcher(None, normalize(left), normalize(right)).ratio() * 100, 1)


def normalize_position_grade(position):
    """Split a position into its normalized base, grade type and numbers."""
    match = GRADE_PATTERN.match(str(position or "").strip())
    if not match:
        return None

    grade_type = match.group("type").casefold()
    normalized_type = "разряд" if grade_type.startswith("разряд") else "категория"
    base = " ".join(normalize(match.group("base")).split())
    numbers = tuple(
        int(number.strip())
        for number in match.group("numbers").split(",")
    )
    if not base or not numbers:
        return None

    return {
        "Основа": base,
        "Тип": normalized_type,
        "Номера": numbers,
    }


def find_grade_matches(position_pairs, norm_pairs, unmatched_position_keys):
    norms_by_department = {}
    for norm_key, norm_row in norm_pairs.items():
        parsed_norm = normalize_position_grade(norm_row["Должность"])
        if not parsed_norm or len(parsed_norm["Номера"]) < 2:
            continue
        norms_by_department.setdefault(norm_key[0], []).append(
            (norm_key, norm_row, parsed_norm)
        )

    grade_matches = []
    covered_position_keys = set()

    for position_key in sorted(unmatched_position_keys):
        position_row = position_pairs[position_key]
        parsed_position = normalize_position_grade(position_row["Должность"])
        if not parsed_position or len(parsed_position["Номера"]) != 1:
            continue

        position_number = parsed_position["Номера"][0]
        for _norm_key, norm_row, parsed_norm in norms_by_department.get(
            position_key[0],
            [],
        ):
            if parsed_position["Основа"] != parsed_norm["Основа"]:
                continue
            if parsed_position["Тип"] != parsed_norm["Тип"]:
                continue
            if position_number not in parsed_norm["Номера"]:
                continue

            covered_position_keys.add(position_key)
            grade_matches.append(
                {
                    "Подразделение": position_row["Подразделение"],
                    "Должность из штатки": position_row["Должность"],
                    "Групповая должность из норм": norm_row["Должность"],
                    "Тип": parsed_position["Тип"],
                    "Номер из штатки": position_number,
                    "Номера из нормы": ", ".join(
                        str(number)
                        for number in parsed_norm["Номера"]
                    ),
                    "Комментарий": GRADE_MATCH_COMMENT,
                }
            )

    return grade_matches, covered_position_keys


def find_possible_matches(missing_norm_rows, missing_position_rows):
    possible_matches = []

    for position_row in missing_norm_rows:
        for norm_row in missing_position_rows:
            department_similarity = similarity_percent(
                position_row["Подразделение"],
                norm_row["Подразделение"],
            )
            position_similarity = similarity_percent(
                position_row["Должность"],
                norm_row["Должность"],
            )

            if position_similarity < 70 and not (
                department_similarity >= 70 and position_similarity >= 50
            ):
                continue

            possible_matches.append(
                {
                    "Подразделение из штатки": position_row["Подразделение"],
                    "Должность из штатки": position_row["Должность"],
                    "Похожее подразделение из норм": norm_row["Подразделение"],
                    "Похожая должность из норм": norm_row["Должность"],
                    "Сходство подразделения, %": department_similarity,
                    "Сходство должности, %": position_similarity,
                    "Комментарий": POSSIBLE_MATCH_COMMENT,
                }
            )

    return sorted(
        possible_matches,
        key=lambda row: (
            normalize(row["Подразделение из штатки"]),
            normalize(row["Должность из штатки"]),
            -row["Сходство должности, %"],
            -row["Сходство подразделения, %"],
        ),
    )


def append_possible_match_rows(sheet, rows):
    for row in rows:
        sheet.append(
            [
                row["Подразделение из штатки"],
                row["Должность из штатки"],
                row["Похожее подразделение из норм"],
                row["Похожая должность из норм"],
                row["Сходство подразделения, %"],
                row["Сходство должности, %"],
                row["Комментарий"],
            ]
        )


def append_grade_match_rows(sheet, rows):
    for row in rows:
        sheet.append(
            [
                row["Подразделение"],
                row["Должность из штатки"],
                row["Групповая должность из норм"],
                row["Тип"],
                row["Номер из штатки"],
                row["Номера из нормы"],
                row["Комментарий"],
            ]
        )


def create_report(
    matched_rows,
    missing_norm_rows,
    missing_position_rows,
    possible_match_rows,
    grade_match_rows,
):
    workbook = Workbook()
    workbook.remove(workbook.active)

    sheets = [
        ("Есть в штатке и в нормах", matched_rows),
        ("Есть в штатке, нет в нормах", missing_norm_rows),
        ("Есть в нормах, нет в штатке", missing_position_rows),
    ]

    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append(POSITIONS_COLUMNS)
        append_report_rows(sheet, rows)
        style_header(sheet)
        set_widths(sheet)
        sheet.freeze_panes = "A2"

    possible_match_sheet = workbook.create_sheet("Возможные совпадения")
    possible_match_sheet.append(
        [
            "Подразделение из штатки",
            "Должность из штатки",
            "Похожее подразделение из норм",
            "Похожая должность из норм",
            "Сходство подразделения, %",
            "Сходство должности, %",
            "Комментарий",
        ]
    )
    append_possible_match_rows(possible_match_sheet, possible_match_rows)
    style_header(possible_match_sheet)
    set_possible_matches_widths(possible_match_sheet)
    possible_match_sheet.freeze_panes = "A2"
    for row in possible_match_sheet.iter_rows(
        min_row=2,
        min_col=5,
        max_col=6,
    ):
        for cell in row:
            cell.number_format = "0.0"

    grade_match_sheet = workbook.create_sheet("Совпадения по разрядам")
    grade_match_sheet.append(
        [
            "Подразделение",
            "Должность из штатки",
            "Групповая должность из норм",
            "Тип",
            "Номер из штатки",
            "Номера из нормы",
            "Комментарий",
        ]
    )
    append_grade_match_rows(grade_match_sheet, grade_match_rows)
    style_header(grade_match_sheet)
    set_grade_matches_widths(grade_match_sheet)
    grade_match_sheet.freeze_panes = "A2"

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook.save(REPORT_PATH)


def main():
    if not POSITIONS_PATH.exists():
        create_positions_placeholder(POSITIONS_PATH)
        print(f"Создан учебный файл-заглушка: {POSITIONS_PATH}")

    position_rows = read_table(POSITIONS_PATH, POSITIONS_COLUMNS)
    norm_rows = read_table(NORMS_PATH, NORM_COLUMNS)

    position_pairs = unique_pairs(position_rows, include_extra=True)
    norm_pairs = unique_pairs(norm_rows)

    matched_rows = [
        position_pairs[key]
        for key in sorted(position_pairs.keys() & norm_pairs.keys())
    ]
    unmatched_position_keys = position_pairs.keys() - norm_pairs.keys()
    grade_match_rows, covered_position_keys = find_grade_matches(
        position_pairs,
        norm_pairs,
        unmatched_position_keys,
    )
    missing_norm_rows = [
        position_pairs[key]
        for key in sorted(unmatched_position_keys - covered_position_keys)
    ]
    missing_position_rows = [
        norm_pairs[key]
        for key in sorted(norm_pairs.keys() - position_pairs.keys())
    ]
    possible_match_rows = find_possible_matches(
        missing_norm_rows,
        missing_position_rows,
    )

    create_report(
        matched_rows,
        missing_norm_rows,
        missing_position_rows,
        possible_match_rows,
        grade_match_rows,
    )

    print(f"Отчет создан: {REPORT_PATH}")
    print(f"Есть в штатке и в нормах: {len(matched_rows)}")
    print(f"Есть в штатке, нет в нормах: {len(missing_norm_rows)}")
    print(f"Есть в нормах, нет в штатке: {len(missing_position_rows)}")
    print(f"Возможные совпадения: {len(possible_match_rows)}")
    print(f"Совпадения по разрядам: {len(grade_match_rows)}")


if __name__ == "__main__":
    main()
