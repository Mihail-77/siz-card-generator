from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


BASE_DIR = Path(__file__).parent


def style_header(sheet):
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def create_norms_file(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Нормы"
    sheet.append(
        [
            "Подразделение",
            "Должность",
            "Наименование СИЗ",
            "Норма выдачи",
            "Срок носки",
            "Основание",
            "Примечание",
        ]
    )
    sheet.append(
        [
            "Производственный участок",
            "Слесарь по ремонту оборудования",
            "Пример позиции СИЗ 1 (замените на реальную норму)",
            "Укажите по утвержденным нормам",
            "Укажите срок",
            "Учебное основание, заменить на утвержденные нормы",
            "Учебная строка, не является нормой",
        ]
    )
    sheet.append(
        [
            "Склад",
            "Кладовщик",
            "Пример позиции СИЗ 2 (замените на реальную норму)",
            "Укажите по утвержденным нормам",
            "Укажите срок",
            "Учебное основание, заменить на утвержденные нормы",
            "Учебная строка, не является нормой",
        ]
    )

    style_header(sheet)
    widths = [28, 30, 46, 30, 18, 42, 34]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    workbook.save(path)


def create_employees_file(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Работники"
    sheet.append(
        [
            "Табельный номер",
            "ФИО",
            "Подразделение",
            "Должность",
            "Дата приема",
            "Пол",
            "Рост",
            "Размер одежды",
            "Размер обуви",
            "Размер головного убора",
            "Размер перчаток",
        ]
    )
    sheet.append(
        [
            "0001",
            "Иванов Иван Иванович",
            "Производственный участок",
            "Слесарь по ремонту оборудования",
            "2024-01-15",
            "Учебно: мужской",
            "Учебно: 180",
            "Учебно: 52",
            "Учебно: 42",
            "Учебно: 58",
            "Учебно: 10",
        ]
    )
    sheet.append(
        [
            "0002",
            "Петров Петр Петрович",
            "Склад",
            "Кладовщик",
            "2023-08-01",
            "Учебно: мужской",
            "Учебно: 176",
            "Учебно: 50",
            "Учебно: 41",
            "Учебно: 57",
            "Учебно: 9",
        ]
    )
    sheet.append(
        [
            "0003",
            "Сидорова Анна Сергеевна",
            "Бухгалтерия",
            "Бухгалтер",
            "2022-05-10",
            "Учебно: женский",
            "Учебно: 168",
            "Учебно: 46",
            "Учебно: 38",
            "Учебно: 56",
            "Учебно: 7",
        ]
    )

    style_header(sheet)
    widths = [18, 30, 30, 34, 16, 18, 16, 18, 16, 24, 18]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    workbook.save(path)


def create_template(path):
    workbook = Workbook()
    front_sheet = workbook.active
    front_sheet.title = "Лицевая сторона"
    back_sheet = workbook.create_sheet("Оборотная сторона")

    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")

    front_widths = [30, 24, 14, 14, 18, 22, 14, 14, 20, 24, 18]
    back_widths = [30, 24, 14, 14, 18, 22, 14, 14, 20, 24]
    for index, width in enumerate(front_widths, start=1):
        front_sheet.column_dimensions[chr(64 + index)].width = width
    for index, width in enumerate(back_widths, start=1):
        back_sheet.column_dimensions[chr(64 + index)].width = width

    for sheet in [front_sheet, back_sheet]:
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = 9
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.4
        sheet.page_margins.right = 0.4
        sheet.page_margins.top = 0.6
        sheet.page_margins.bottom = 0.4

    front_sheet.merge_cells("H1:K1")
    front_sheet["H1"] = "Приложение № 2"
    front_sheet["H1"].alignment = Alignment(horizontal="center")
    front_sheet.merge_cells("H2:K5")
    front_sheet["H2"] = (
        "к Правилам обеспечения работников средствами индивидуальной защиты "
        "и смывающими средствами"
    )
    front_sheet["H2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    front_sheet.merge_cells("H7:K7")
    front_sheet["H7"] = "Рекомендуемый образец"
    front_sheet["H7"].alignment = Alignment(horizontal="right")

    front_sheet.merge_cells("A6:K6")
    front_sheet["A6"] = "Личная карточка учета выдачи СИЗ"
    front_sheet["A6"].font = Font(bold=True, size=14)
    front_sheet["A6"].alignment = Alignment(horizontal="center")
    front_sheet.merge_cells("A8:K8")
    front_sheet["A8"] = "ЛИЧНАЯ КАРТОЧКА №"
    front_sheet["A8"].font = Font(bold=True, size=13)
    front_sheet["A8"].alignment = Alignment(horizontal="center")
    front_sheet.merge_cells("A9:K9")
    front_sheet["A9"] = "учета выдачи СИЗ"
    front_sheet["A9"].alignment = Alignment(horizontal="center")

    front_sheet.merge_cells("A10:F10")
    front_sheet["A10"] = "Сведения о работнике"
    front_sheet.merge_cells("H10:K10")
    front_sheet["H10"] = "Размеры"
    for cell in ["A10", "H10"]:
        front_sheet[cell].font = Font(bold=True)
        front_sheet[cell].fill = header_fill
        front_sheet[cell].alignment = Alignment(horizontal="center")

    worker_rows = [
        (11, "Фамилия"),
        (12, "Имя"),
        (13, "Табельный номер"),
        (14, "Структурное подразделение"),
        (15, "Профессия (должность)"),
        (16, "Дата поступления на работу"),
        (17, "Дата изменения профессии (должности) или"),
        (18, "перевода в другое структурное подразделение"),
    ]
    for row_number, label in worker_rows:
        front_sheet.cell(row=row_number, column=1, value=label)
        front_sheet.cell(row=row_number, column=1).font = Font(bold=True)
    front_sheet["E12"] = "Отчество (при наличии)"
    front_sheet["E12"].font = Font(bold=True)

    size_rows = [
        (11, "Пол"),
        (12, "Рост"),
        (13, "Размер:"),
        (14, "одежды"),
        (15, "обуви"),
        (16, "головного убора"),
        (17, "СИЗОД"),
        (18, "СИЗ рук"),
    ]
    for row_number, label in size_rows:
        front_sheet.cell(row=row_number, column=8, value=label)
        front_sheet.cell(row=row_number, column=8).font = Font(bold=True)

    front_sheet.merge_cells("A21:K21")
    front_sheet["A21"] = "Положенные работнику СИЗ"
    front_sheet["A21"].font = Font(bold=True)
    front_sheet["A21"].alignment = Alignment(horizontal="center")
    front_sheet["A21"].fill = header_fill

    front_headers = [
        "Наименование СИЗ",
        "Пункт Норм / основание",
        "Единица измерения, периодичность выдачи",
        "Количество на период",
        "Примечание",
    ]
    for column, header in enumerate(front_headers, start=1):
        cell = front_sheet.cell(row=22, column=column, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in front_sheet.iter_rows(min_row=23, max_row=34, min_col=1, max_col=5):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    front_sheet["A36"] = "Ответственное лицо за ведение карточек учета выдачи СИЗ"
    front_sheet["C37"] = "(подпись)"
    front_sheet["F37"] = "(фамилия, инициалы)"
    front_sheet["C36"].border = Border(bottom=thin)
    front_sheet["D36"].border = Border(bottom=thin)
    front_sheet["F36"].border = Border(bottom=thin)
    front_sheet["G36"].border = Border(bottom=thin)

    back_sheet.merge_cells("A1:J1")
    back_sheet["A1"] = "Оборотная сторона личной карточки"
    back_sheet["A1"].font = Font(bold=True)
    back_sheet["A1"].alignment = Alignment(horizontal="center")

    back_sheet.merge_cells("A3:A4")
    back_sheet["A3"] = "Наименование СИЗ"
    back_sheet.merge_cells("B3:B4")
    back_sheet["B3"] = "Модель, марка, артикул, класс защиты СИЗ, дерматологических СИЗ"
    back_sheet.merge_cells("C3:F3")
    back_sheet["C3"] = "Выдано"
    back_sheet.merge_cells("G3:J3")
    back_sheet["G3"] = "Возвращено"

    issue_headers = [
        "дата",
        "количество",
        "Лично /\nкарточка",
        "подпись получившего СИЗ",
        "дата",
        "количество",
        "подпись сдавшего СИЗ",
        "Акт списания, дата и номер",
    ]
    for column, header in enumerate(issue_headers, start=3):
        back_sheet.cell(row=4, column=column, value=header)
    for column in range(1, 11):
        back_sheet.cell(row=5, column=column, value=column)

    for row in back_sheet.iter_rows(min_row=3, max_row=20, min_col=1, max_col=10):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number in range(3, 6):
        for cell in back_sheet[row_number]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

    back_sheet["A23"] = "* информация указывается только для дерматологических СИЗ"
    back_sheet["A24"] = (
        "** информация указывается для всех СИЗ, кроме дерматологических СИЗ "
        "и СИЗ однократного применения"
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main():
    create_norms_file(BASE_DIR / "norms.xlsx")
    create_employees_file(BASE_DIR / "employees.xlsx")
    create_template(BASE_DIR / "templates" / "card_template.xlsx")
    print("Созданы файлы norms.xlsx, employees.xlsx и templates/card_template.xlsx")


if __name__ == "__main__":
    main()
